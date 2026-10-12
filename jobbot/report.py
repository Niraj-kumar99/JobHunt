"""Writes the results: an HTML report page, CSV and Excel files, and a GitHub run summary."""
from __future__ import annotations

import csv
import html
import os
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

MATCH_COLS = [("company", "Company"), ("title", "Job title"), ("location", "Location"),
              ("posted", "Posted"), ("experience", "Experience"), ("hq_country", "HQ country"),
              ("sector", "Sector"), ("url", "Link"), ("found_on", "Found on"), ("notes", "Notes")]
COV_COLS = [("company", "Company"), ("country", "Country"), ("platform", "Platform"),
            ("status", "Status"), ("jobs_read", "Jobs read"), ("title_matches", "Title matches"),
            ("new_checked", "New jobs checked"), ("new_matches", "New matches"),
            ("filtered_out", "Filtered out (reason: count)"), ("message", "Message"),
            ("platform_id", "Platform ID")]


def _posted(m: dict) -> str:
    return f"{m['posted']} (first seen)" if m.get("date_is_first_seen") else m.get("posted", "")


# ----------------------------------------------------------------------------- spreadsheets
def _sheet(ws, cols, rows, link_key=None):
    ws.append([label for _, label in cols])
    for r in rows:
        ws.append([_posted(r) if k == "posted" else r.get(k, "") for k, _ in cols])
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial")
    if link_key:
        idx = [k for k, _ in cols].index(link_key) + 1
        for i in range(2, ws.max_row + 1):
            cell = ws.cell(row=i, column=idx)
            if cell.value:
                cell.hyperlink = cell.value
                cell.value = "Open job"
                cell.font = Font(name="Arial", color="0563C1", underline="single")
    widths = {"Company": 26, "Job title": 46, "Location": 30, "Posted": 20, "Experience": 22,
              "Link": 12, "Message": 50, "Filtered out (reason: count)": 50, "Notes": 30}
    for i, (_, label) in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(label, 16)
    ws.freeze_panes = "A2"
    if rows:
        ref = f"A1:{ws.cell(row=ws.max_row, column=len(cols)).coordinate}"
        t = Table(displayName=ws.title.replace(" ", ""), ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)


def write_xlsx(path: Path, new, history, coverage):
    wb = Workbook()
    ws = wb.active
    ws.title = "New matches"
    _sheet(ws, MATCH_COLS, new, "url")
    _sheet(wb.create_sheet("Recent matches"), MATCH_COLS, history, "url")
    _sheet(wb.create_sheet("Coverage"), COV_COLS, coverage)
    wb.save(path)


def write_csv(path: Path, cols, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([label for _, label in cols])
        for r in rows:
            w.writerow([_posted(r) if k == "posted" else r.get(k, "") for k, _ in cols])


# ----------------------------------------------------------------------------- HTML page
CSS = """
:root{--bg:#f7f7f5;--card:#fff;--ink:#1c1c1a;--muted:#6b6b66;--line:#e4e3de;--accent:#1f4fd6;--new:#e8f0ff;--warn:#a15c00}
@media (prefers-color-scheme:dark){:root{--bg:#141413;--card:#1d1d1b;--ink:#ecebe6;--muted:#9b9a94;--line:#2e2e2b;--accent:#7aa2ff;--new:#1c2640;--warn:#e0a34d}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 system-ui,-apple-system,Segoe UI,Roboto,sans-serif}
main{max-width:1100px;margin:0 auto;padding:28px 16px 60px}h1{font-size:26px;margin:0 0 4px}h2{font-size:18px;margin:34px 0 10px}
.sub{color:var(--muted);margin:0 0 20px}.stats{display:flex;flex-wrap:wrap;gap:10px;margin:18px 0}
.stat{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:10px 14px;min-width:130px}
.stat b{display:block;font-size:22px}.stat span{color:var(--muted);font-size:13px}
.card{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px 16px;margin:10px 0}
.card.new{background:var(--new)}.card h3{margin:0;font-size:16px}.card h3 a{color:var(--accent);text-decoration:none}
.card h3 a:hover{text-decoration:underline}.meta{color:var(--muted);font-size:13.5px;margin-top:4px}
.tag{display:inline-block;border:1px solid var(--line);border-radius:999px;padding:0 8px;margin-right:6px;font-size:12px}
.empty{color:var(--muted);padding:12px 0}input{width:100%;padding:9px 12px;border:1px solid var(--line);border-radius:8px;
background:var(--card);color:var(--ink);font:inherit;margin:4px 0 6px}details{margin-top:8px}summary{cursor:pointer;color:var(--muted)}
table{width:100%;border-collapse:collapse;font-size:13px;margin-top:8px}td,th{border-bottom:1px solid var(--line);padding:6px 4px;text-align:left;vertical-align:top}
.err{color:var(--warn)}.wrap{overflow-x:auto}
"""


def _card(m: dict, is_new: bool) -> str:
    e = html.escape
    tags = f'<span class="tag">{e(m.get("experience", ""))}</span>' if m.get("experience") else ""
    notes = f' · {e(m["notes"])}' if m.get("notes") else ""
    return (f'<div class="card{" new" if is_new else ""}" data-q="{e((m["company"] + " " + m["title"] + " " + m.get("location", "")).lower())}">'
            f'<h3><a href="{e(m["url"])}" target="_blank" rel="noopener">{e(m["title"])}</a></h3>'
            f'<div class="meta"><b>{e(m["company"])}</b> · {e(m.get("location") or "location not stated")} · '
            f'posted {e(_posted(m))}{notes}</div><div class="meta">{tags}</div></div>')


def write_html(path: Path, new, history, coverage, info):
    e = html.escape
    ok = sum(1 for c in coverage if c["status"] == "ok")
    platforms = Counter(c["platform"] or "unknown" for c in coverage if c["status"] == "ok")
    bad = [c for c in coverage if c["status"] != "ok"]
    when = info["finished"].strftime("%d %b %Y, %H:%M UTC")
    cfg = info["config"]
    first = ('<p class="sub"><b>First run:</b> this run built the list of existing jobs. Only postings from the last '
             f'{cfg["max_age_days"]} days were reported; from now on you will see only new ones.</p>') if info["first_run"] else ""
    new_keys = {m["key"] for m in new}
    older = [m for m in history if m["key"] not in new_keys]
    new_html = "".join(_card(m, True) for m in new) or '<div class="empty">No new matching jobs in this run.</div>'
    old_html = "".join(_card(m, False) for m in older) or '<div class="empty">Nothing yet.</div>'
    plat = ", ".join(f"{k} {v}" for k, v in platforms.most_common())
    rows = "".join(
        f'<tr><td>{e(c["company"])}</td><td>{e(c["platform"])}</td><td class="err">{e(c["status"])}</td>'
        f'<td>{e(c["message"])}</td></tr>' for c in bad)
    page = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>Job Matches</title><style>{CSS}</style></head>
<body><main>
<h1>Entry-level PM &amp; BA jobs in Europe</h1>
<p class="sub">Last run {when} · {info['companies']} companies · posted within {cfg['max_age_days']} days ·
max {cfg['max_experience_years']} yr experience · English · Europe</p>
{first}
<div class="stats">
<div class="stat"><b>{len(new)}</b><span>new matches</span></div>
<div class="stat"><b>{len(older)}</b><span>earlier matches ({cfg['history_days']} days)</span></div>
<div class="stat"><b>{ok}/{len(coverage)}</b><span>companies scanned</span></div>
</div>
<input id="q" placeholder="Filter by company, title or city…" oninput="f(this.value)">
<h2>New in this run</h2>{new_html}
<h2>Earlier matches</h2>{old_html}
<h2>Coverage</h2>
<p class="sub">Read via: {e(plat)}.</p>
<details><summary>{len(bad)} companies couldn't be read this run</summary>
<p class="sub">Usually the careers site needs a browser or blocked the request. You can set the job platform for these in companies.csv (see README).</p>
<div class="wrap"><table><tr><th>Company</th><th>Platform</th><th>Status</th><th>Message</th></tr>{rows}</table></div></details>
</main>
<script>function f(v){{v=v.toLowerCase();document.querySelectorAll('.card').forEach(c=>{{c.style.display=c.dataset.q.includes(v)?'':'none'}})}}</script>
</body></html>"""
    path.write_text(page, encoding="utf-8")


# ----------------------------------------------------------------------------- GitHub summary
def write_summary(new, coverage, info):
    target = os.environ.get("GITHUB_STEP_SUMMARY")
    if not target:
        return
    ok = sum(1 for c in coverage if c["status"] == "ok")
    lines = [f"## {len(new)} new matching jobs", "",
             f"{ok}/{len(coverage)} companies scanned. Download **job-matches** below for the spreadsheet.", ""]
    if info["first_run"]:
        lines += ["_First run: existing jobs were recorded; future runs show only new ones._", ""]
    if new:
        lines += ["| Company | Job | Location | Posted | Experience |", "|---|---|---|---|---|"]
        for m in new[:200]:
            title = m["title"].replace("|", "/")
            lines.append(f"| {m['company']} | [{title}]({m['url']}) | {m['location'].replace('|', '/')} | "
                         f"{_posted(m)} | {m['experience']} |")
    with open(target, "a", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")


def write_all(root: Path, new, history, coverage, info):
    out, docs = root / "output", root / "docs"
    out.mkdir(exist_ok=True)
    docs.mkdir(exist_ok=True)
    write_csv(out / "new_matches.csv", MATCH_COLS, new)
    write_csv(out / "coverage.csv", COV_COLS, coverage)
    write_xlsx(out / "job_matches.xlsx", new, history, coverage)
    write_html(docs / "index.html", new, history, coverage, info)
    write_summary(new, coverage, info)
